"""T0.3 手工积木原型：HN（hn.algolia.com）搜索 → 行数据 → Excel 报告。

锚：docs/07-MS0执行计划.md T0.3a Interfaces。
SELECTORS 固化自 T0.2b 结构报告（spikes/evidence/t0.2_*.json，2026-10-14）。
"""
import json
import re
from pathlib import Path
from urllib.parse import quote

from openpyxl import Workbook, load_workbook

from common.browser import new_page
from site_inspect.inspect import _goto_with_retry

SEARCH_URL_TMPL = "https://hn.algolia.com/?q={query}"
ROW_FIELDS = ("position", "title", "url", "points", "num_comments")

SELECTORS = {
    "result_item": "article.Story",
}
# 条目内部用相对选择器（链式 locator 以 item 为根，全局路径会匹配不到）
ITEM_SELECTORS = {
    "title": ".Story_title a",
    "link": ".Story_link",
    "meta": ".Story_meta",
}


def _meta_int(meta_text: str, pattern: str) -> int:
    m = re.search(pattern, meta_text or "")
    return int(m.group(1)) if m else 0


def _extract_row(page, index: int) -> dict:
    item = page.locator(SELECTORS["result_item"]).nth(index)
    title_loc = item.locator(ITEM_SELECTORS["title"]).first
    title = title_loc.inner_text().strip()
    link = item.locator(ITEM_SELECTORS["link"])
    # 外链缺失（Ask HN 类纯文本帖）时回落到 HN item 链接
    url = (link if link.count() > 0 else title_loc).first.get_attribute("href") or ""
    meta = item.locator(ITEM_SELECTORS["meta"]).first.inner_text()
    return {
        "position": index + 1,
        "title": title,
        "url": url,
        "points": _meta_int(meta, r"(\d+)\s*points?"),
        "num_comments": _meta_int(meta, r"(\d+)\s*comments?"),
    }


def run_search(query: str, out_xlsx, out_json, limit: int = 10) -> list[dict]:
    """搜索并提取前 limit 条结果，写 JSON + xlsx（回读校验后返回 rows）。"""
    with new_page() as page:
        _goto_with_retry(
            page,
            SEARCH_URL_TMPL.format(query=quote(query)),
            wait_until="domcontentloaded",
            timeout=30_000,
        )
        page.wait_for_function(
            "arg => document.querySelectorAll(arg.sel).length >= arg.limit",
            arg={"sel": SELECTORS["result_item"], "limit": limit},
            timeout=20_000,
        )
        items = page.locator(SELECTORS["result_item"])
        if items.count() < limit:
            raise RuntimeError(
                f"selector_missing: expected >={limit}, got {items.count()}"
            )
        rows = [_extract_row(page, i) for i in range(limit)]

    for row in rows:
        assert row["title"] and row["url"].startswith("http"), f"empty row fields: {row}"

    out_json, out_xlsx = Path(out_json), Path(out_xlsx)
    out_json.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "report"
    ws.append(list(ROW_FIELDS))
    for row in rows:
        ws.append([row[f] for f in ROW_FIELDS])
    wb.save(out_xlsx)

    check = load_workbook(out_xlsx)["report"]
    assert [c.value for c in check[1]] == list(ROW_FIELDS)
    assert check.max_row == len(rows) + 1
    return rows
