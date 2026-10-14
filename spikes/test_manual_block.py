"""T0.3 验收测试：手工积木原型——标准站点搜索 → 行数据 → Excel 报告。

锚：docs/07-MS0执行计划.md T0.3a；docs/05-开发计划.md §4 阶段 0 T0.3。
查询词用 python：patchright 全站仅 5 条 < limit=10（裁定见计划台账 2026-10-14）。
"""
import json

from openpyxl import load_workbook

from manual_block.hn_search_block import ROW_FIELDS, run_search


def test_search_to_excel_report(tmp_path):
    rows = run_search(
        "python", out_xlsx=tmp_path / "report.xlsx", out_json=tmp_path / "rows.json"
    )
    assert len(rows) == 10
    assert all(r["title"] and r["url"].startswith("http") for r in rows)
    assert [r["position"] for r in rows] == list(range(1, 11))
    ws = load_workbook(tmp_path / "report.xlsx")["report"]  # 损坏/错位 → 在此失败
    assert [c.value for c in ws[1]] == list(ROW_FIELDS)
    assert ws.max_row == 11
    assert ws["A2"].value == 1
    assert json.loads((tmp_path / "rows.json").read_text(encoding="utf-8")) == rows
